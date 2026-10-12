"""CSV / Excel / PDF exports of tabular data (INV-08, REP-06, section 5.2)."""
import csv
import io
from collections.abc import Sequence
from typing import Any

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from ..audit import log_event

Columns = Sequence[tuple[str, str]]  # (key, header)


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return v


def to_csv(rows: list[dict], columns: Columns) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for _, h in columns])
    for r in rows:
        w.writerow([_cell(r.get(k)) for k, _ in columns])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens UTF-8 (Arabic) correctly


def to_xlsx(rows: list[dict], columns: Columns, title: str = "Export") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Export"
    ws.append([h for _, h in columns])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([_cell(r.get(k)) for k, _ in columns])
    for i, (_, h) in enumerate(columns, start=1):
        width = max([len(str(h))] + [len(str(_cell(r.get(columns[i - 1][0])))) for r in rows[:500]])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(60, width + 2)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def file_response(content: bytes, fmt: str, filename: str, inline: bool = False) -> Response:
    disp = "inline" if inline else "attachment"
    return Response(content, media_type=MEDIA[fmt],
                    headers={"Content-Disposition": f'{disp}; filename="{filename}.{fmt}"'})


def export(db: Session, rows: list[dict], columns: Columns, fmt: str, name: str, entity: str,
           title: str | None = None, summary: list[tuple[str, str]] | None = None) -> Response:
    if fmt == "csv":
        content = to_csv(rows, columns)
    elif fmt == "xlsx":
        content = to_xlsx(rows, columns, title or name)
    elif fmt == "pdf":
        from .pdf import table_pdf

        content = table_pdf(db, title or name, columns, rows, summary or [])
    else:
        from ..deps import http_error

        raise http_error(422, "format must be csv, xlsx or pdf", "validation")
    log_event(db, "export", entity, None, after={"format": fmt, "rows": len(rows), "name": name})
    db.commit()
    return file_response(content, fmt, name)
